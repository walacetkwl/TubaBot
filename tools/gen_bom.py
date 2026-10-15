#!/usr/bin/env python3
"""Gera docs/lista_de_materiais.xlsx (lista de compras da v3, ES3C28P). Requer openpyxl.
Uso: python3 tools/gen_bom.py docs/lista_de_materiais.xlsx
Para gravar os totais calculados no arquivo (útil em visualizadores), abra e salve no LibreOffice/Excel."""
import sys, urllib.parse as up
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

out = sys.argv[1]
def ml(q): return "https://lista.mercadolivre.com.br/" + q.replace(" ", "-") + "_OrderId_PRICE_NoIndex_True"
def ml_rel(q): return "https://lista.mercadolivre.com.br/" + q.replace(" ", "-")
def sh(q): return "https://shopee.com.br/search?keyword=" + up.quote(q) + "&sortBy=sales"

H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F3A68")
SEC = PatternFill("solid", fgColor="E8EEF8"); OPT = PatternFill("solid", fgColor="FFF6E0")
LINK = Font(color="1155CC", underline="single")
thin = Side(style="thin", color="C8D0DC"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
ws = wb.active; ws.title = "Lista de compras"
ws["A1"] = "TubaBot v3 (ES3C28P) — lista de materiais por unidade (Brasil)"; ws["A1"].font = Font(bold=True, size=15, color="1F3A68")
ws["A2"] = ("Preços em R$ são FAIXAS ESTIMADAS (out/2026) a partir de preços internacionais pesquisados. "
            "As páginas do Mercado Livre e da Shopee não puderam ser abertas deste ambiente: os links levam a buscas já "
            "filtradas por 'menor preço' e por 'mais vendidos/relevantes' (os anúncios com mais avaliações). Confira no anúncio: "
            "nota ≥ 4,5, mais de 50 vendas e a especificação da coluna 'Conferir no anúncio'. Lite = sem bateria; Pro = com bateria.")
ws["A2"].alignment = WR; ws.merge_cells("A2:M2"); ws.row_dimensions[2].height = 62

cols = ["#", "Item", "Qtd", "Lite / Pro", "Especificação exata", "Conferir no anúncio", "Mercado Livre — menor preço",
        "Mercado Livre — mais relevantes", "Shopee — mais vendidos", "Outra fonte (referência)", "Preço mín. estimado (R$)",
        "Preço máx. estimado (R$)", "Observações"]
r0 = 4
for i, c in enumerate(cols, 1):
    x = ws.cell(r0, i, c); x.font = H; x.fill = HF; x.alignment = Alignment(wrap_text=True, vertical="center"); x.border = B
ws.row_dimensions[r0].height = 34

items = [
 ("Placa ES3C28P — ESP32-S3 + tela IPS 2,8\" 240×320 com TOQUE CAPACITIVO", 1, "Lite e Pro",
  "ESP32-S3R8 (16 MB flash, 8 MB PSRAM), ILI9341V, toque FT6336G, codec ES8311 + amplificador + microfone, conector de alto-falante, conector de bateria com carregador TP4054, USB-C",
  "Modelo exato ES3C28P (o 'C' = capacitivo). NÃO comprar ES3N28P (sem toque) nem a 'ESP32-S3-Touch-LCD-2.8' da Waveshare (outra placa). A foto deve mostrar os conectores SPEAKER e BAT.",
  "es3c28p esp32-s3 2.8", "esp32-s3 2.8 tela touch capacitiva es3c28p", "es3c28p esp32 s3 2.8",
  ("AliExpress — busca ES3C28P", "https://pt.aliexpress.com/w/wholesale-ES3C28P.html"),
  90, 180, "Referência internacional ~US$ 15 (fabricante QDtech/LCDWIKI). No AliExpress com o Remessa Conforme costuma sair mais barato que no ML."),
 ("Alto-falante 8 Ω 1,5–2 W (ou 4 Ω 2 W), pequeno", 1, "Lite e Pro",
  "Cápsula 28–40 mm ou retangular 20×30 mm; 8 Ω 1,5 W ou 4 Ω 2 W (valores da LCDWIKI)",
  "Impedância e potência; já com fios. Confira o plugue do conector SPEAKER da sua placa (MX1.25 ou PH2.0) — se não bater, use o kit de rabichos.",
  "alto falante 8 ohms 2w 28mm", "mini alto falante 8 ohms 1.5w", "alto falante 8ohm 2w 28mm",
  ("Busca geral ML (alto-falante 8 Ω 2 W)", "https://lista.mercadolivre.com.br/alto-falante-8-ohms-2w"),
  8, 25, "Kits com 2–5 unidades costumam custar o mesmo que uma."),
 ("Bateria LiPo 3,7 V 1000–1500 mAh", 1, "só Pro",
  "LiPo 1S 3,7 V com proteção (PCM); 1000 a 1500 mAh; espessura conforme o corpo (ex.: 103040, 803450)",
  "COM proteção; tamanho que caiba no corpo; plugue compatível com o conector BAT (MX1.25 ou PH2.0). Antes de plugar, confira a polaridade com o multímetro.",
  "bateria lipo 3.7v 1000mah", "bateria lipo 3.7v 1200mah com protecao", "bateria lipo 3.7v 1000mah",
  ("Busca geral ML (bateria 3,7 V 1000 mAh)", "https://lista.mercadolivre.com.br/bateria-3.7v-1000mah"),
  25, 55, "O carregador da placa (TP4054) carrega a até ~0,5 A: 1000–1500 mAh carregam em 2–4 h."),
 ("Kit de rabichos MX1.25 e PH2.0 de 2 vias", 1, "Lite e Pro",
  "Conectores 2 vias com fio (passo 1,25 mm e 2,0 mm)",
  "Um kit com os dois passos resolve qualquer diferença entre placa, alto-falante e bateria.",
  "kit conector mx1.25 ph2.0 2 pinos com fio", "cabo jst ph2.0 2 vias", "mx1.25 ph2.0 2pin cabo kit",
  None, 12, 35, "Exige solda simples (o ferro da v1 serve). Serve para muitas unidades."),
 ("Filamento PLA 1,75 mm na cor do corpo (ou impressão sob encomenda)", 1, "Lite e Pro",
  "PLA 1 kg nas cores do produto: vermelho, azul, branco ou grafite/preto",
  "PLA de marca com boas avaliações, na cor da variante. Serviço: envie o STL do corpo do TubaBot.",
  "filamento pla 1.75mm 1kg", "filamento pla 1kg vermelho", "filamento pla 1.75mm 1kg",
  ("Filamento PLA — busca ML", "https://lista.mercadolivre.com.br/filamento-pla-1.75mm-1kg_OrderId_PRICE_NoIndex_True"),
  15, 120, "Um corpo usa ~40–80 g: 1 kg rende vários TubaBots (mínimo = custo por unidade)."),
 ("Parafusos M2 × 6 mm (ou M2,5) + porcas/insertos", 4, "Lite e Pro",
  "Para os 4 furos da placa ES3C28P",
  "Confira o diâmetro dos furos da sua placa (M2 ou M2,5).",
  "parafuso m2 6mm kit", "kit parafuso m2 m2.5 eletronica", "parafuso m2 6mm",
  None, 0.5, 4, "Kits sortidos M2/M2,5 servem para muitas unidades (preço por parafuso)."),
 ("Fita dupla-face de espuma fina", 1, "só Pro",
  "Fita dupla-face de espuma, 0,8–1 mm, 12 mm",
  "Espessura ≤ 1 mm.",
  "fita dupla face espuma fina 12mm", "fita dupla face 3m fina", "fita dupla face fina espuma",
  None, 2, 20, "Prende a bateria no corpo, longe de cantos vivos."),
]
opt = [
 ("Cabo USB-C de dados (Ugreen US287)", "JÁ TENHO (v1)", "Gravação, alimentação e carga (Pro)."),
 ("Multímetro Kapbom DT-9205A", "JÁ TENHO (v1)", "Conferir a polaridade da bateria."),
 ("Ferro de solda 60 W + estanho", "JÁ TENHO (v1)", "Só para adaptar plugues."),
 ("Ímãs de neodímio 8×2 mm", "JÁ TENHO (v1)", "Opcional: fecho do corpo."),
 ("ESP32-C3, OLED 0,96\", TTP223, buzzer, motor MV50", "NÃO USADOS na v3", "Toda a interação agora é pela tela de toque."),
]
r = r0 + 1
for n, it in enumerate(items, 1):
    name, q, req, spec, check, q1, q2, q3, other, pmin, pmax, obs = it
    vals = [n, name, q, req, spec, check, None, None, None, None, pmin, pmax, obs]
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v); c.alignment = WR; c.border = B
        if req.startswith("só"): c.fill = OPT
    for col, qq, fn, label in ((7, q1, ml, "ML · menor preço"), (8, q2, ml_rel, "ML · relevantes"), (9, q3, sh, "Shopee · mais vendidos")):
        c = ws.cell(r, col, label); c.hyperlink = fn(qq); c.font = LINK
    if other:
        c = ws.cell(r, 10, other[0]); c.hyperlink = other[1]; c.font = LINK
    for col in (11, 12): ws.cell(r, col).number_format = 'R$ #,##0.00'
    ws.row_dimensions[r].height = 96
    r += 1
last = r - 1
def total(label, rows_filter):
    global r
    ws.cell(r, 10, label).font = Font(bold=True)
    for col in (11, 12):
        L = get_column_letter(col)
        terms = "+".join(f"C{rr}*{L}{rr}" for rr in rows_filter)
        c = ws.cell(r, col, f"={terms}"); c.font = Font(bold=True); c.number_format = 'R$ #,##0.00'; c.border = B
    r += 1
pro_rows = list(range(r0 + 1, last + 1))
lite_rows = [rr for rr, it in zip(pro_rows, items) if not it[2].startswith("só")]
total("TOTAL estimado — TubaBot Lite", lite_rows)
total("TOTAL estimado — TubaBot Pro", pro_rows)
r += 1
ws.cell(r, 1, "Peças da v1 (não comprar)").font = Font(bold=True, size=12, color="1F3A68"); r += 1
for i, h in enumerate(["#", "Item", "Situação", "Uso na v3"], 1):
    c = ws.cell(r, i, h); c.font = H; c.fill = HF; c.border = B
r += 1
for n, (a, b, c_) in enumerate(opt, 1):
    for i, v in enumerate([n, a, b, c_], 1):
        c = ws.cell(r, i, v); c.alignment = WR; c.border = B; c.fill = SEC
    r += 1
widths = [4, 34, 6, 11, 40, 46, 18, 18, 20, 30, 13, 13, 40]
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "C5"

w2 = wb.create_sheet("Como escolher o anúncio")
rows = [
 ("Critério", "Por quê"),
 ("Reputação do vendedor verde (ML) / 'Indicado' ou 'Loja oficial' (Shopee)", "Menor chance de produto errado e devolução fácil"),
 ("Nota ≥ 4,5 com 50+ avaliações; leia as avaliações com foto", "Confirma que chega o item da foto"),
 ("Placa: o nome ES3C28P aparece no anúncio/placa (C = toque capacitivo)", "ES3N28P não tem toque; a Waveshare 2.8 é outra placa, com outros pinos"),
 ("Bateria (Pro): 'com proteção' e tamanho que caiba no corpo", "Segurança e encaixe"),
 ("Conectores SPEAKER e BAT: pergunte ao vendedor se são MX1.25 ou PH2.0", "Os dois são comuns nesta família de placas"),
 ("Prefira 'Full' (ML) ou envio nacional", "Chega mais rápido, sem taxa de importação surpresa"),
 ("Para importados: Remessa Conforme ativo", "Imposto já incluso no preço exibido"),
 ("Compare preço + frete; kits de alto-falantes/rabichos/parafusos compensam", "Sobra peça para as próximas unidades"),
]
for i, (a, b) in enumerate(rows, 1):
    w2.cell(i, 1, a).alignment = WR; w2.cell(i, 2, b).alignment = WR
    if i == 1:
        for c in (1, 2): w2.cell(i, c).font = H; w2.cell(i, c).fill = HF
w2.column_dimensions["A"].width = 70; w2.column_dimensions["B"].width = 60

w3 = wb.create_sheet("Fontes e notas")
src = [
 ("Fonte", "Link", "O que foi usado"),
 ("LCDWIKI — 2.8inch ESP32-S3 Display (ES3C28P)", "https://www.lcdwiki.com/2.8inch_ESP32-S3_Display", "especificações e pinos"),
 ("BSP ngttai/esp32_s3_es3c28p", "https://github.com/ngttai/esp32_s3_es3c28p", "mapa de pinos, alto-falante 8 Ω 1,5 W / 4 Ω 2 W, bateria com TP4054"),
 ("Launcher — perfil ES3C28P", "https://github.com/bmorcelli/Launcher/issues/373", "pinos de áudio (I²S) e habilitação do amplificador"),
 ("Bruce firmware — PR ES3C28P", "https://github.com/BruceDevices/firmware/pull/2178", "preço de referência ~US$ 15"),
 ("Alibaba — ES3C28P/ES3N28P", "https://www.alibaba.com/product-detail/ES3C28P-ES3N28P-2-8-inch-ESP32_1601603512242.html", "US$ 12–15 (atacado)"),
 ("Mercado Livre — bateria 3,7 V 1000 mAh", "https://lista.mercadolivre.com.br/bateria-3.7v-1000mah", "faixa de preço de baterias"),
 ("Mercado Livre — alto-falante 8 Ω 2 W", "https://lista.mercadolivre.com.br/alto-falante-8-ohms-2w", "faixa de preço de alto-falantes"),
 ("Nota", "", "Preços em R$ estimados em 15/10/2026 (câmbio ~R$ 5,4–5,6/US$ + frete/imposto para importados). Confirme no anúncio antes de comprar."),
]
for i, row in enumerate(src, 1):
    for j, v in enumerate(row, 1):
        c = w3.cell(i, j, v); c.alignment = WR
        if i == 1: c.font = H; c.fill = HF
        elif j == 2 and v: c.hyperlink = v; c.font = LINK
w3.column_dimensions["A"].width = 40; w3.column_dimensions["B"].width = 70; w3.column_dimensions["C"].width = 60
wb.save(out); print("ok", out)
